#!/usr/bin/env python3
"""Refresh the public Samsung dashboard data from the local Excel workbooks.

Source workbooks stay outside this project directory and must not be committed.
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE_DIR = ROOT.parent
MAIN_BOOK = SOURCE_DIR / "Мотивации_26_SM.xlsx"
ANNUAL_BOOK = SOURCE_DIR / "Мотивации SAMSUNG 2026_СВОД.xlsx"
OUT = ROOT / "site-data.js"

STORES = sorted([
    "SMSG_БРН_ГЭЛАКСИ", "SMSG_ИРК_МОДНЫЙ_КВАРТАЛ", "SMSG_ИРК_СИЛЬВЕР_МОЛЛ",
    "SMSG_КМР_ЛАПЛАНДИЯ", "SMSG_НБК_ГАЛЕРЕЯ", "SMSG_НБК_КРАСНЫЙ", "SMSG_НБК_МЕГА",
    "SMSG_НВК_ПЛАНЕТА", "SMSG_ОМС_МЕГА", "SMSG_СГТ_АУРА", "SMSG_СГТ_СИТИ_МОЛЛ",
    "SMSG_ТМК_ИЗУМРУДНЫЙ_ГОРОД", "SMSG_ТМН_ГУДВИН", "SMSG_ТМН_КРИСТАЛЛ",
    "SMSG_ТМН_СИТИ_МОЛЛ", "SMSG_ХБР_БРОСКО_МОЛЛ",
], key=str.casefold)


def clean_number(value, default=0):
    return value if isinstance(value, (int, float)) else default


def date_text(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, str):
        return value.split(":", 1)[-1].strip()
    return ""


def short_period(value, fallback):
    parts = re.findall(r"\d{2}\.\d{2}(?:\.\d{4})?", value)
    if len(parts) >= 2:
        return f"{parts[0][:5]}–{parts[1][:5]}"
    return fallback


def read_team(ws, is_money):
    result = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        name = row[10] if len(row) > 10 else None
        if name not in STORES:
            continue
        plan = clean_number(row[12], None)
        fact = clean_number(row[13], None)
        forecast = clean_number(row[14], None)
        result[name] = {
            "group": clean_number(row[11], None),
            "plan": plan,
            "fact": fact,
            "forecast": forecast,
            "place": clean_number(row[16], 0) if is_money else None,
            "prize": clean_number(row[17], 0) if is_money else None,
        }
    return result


def read_ubear(ws):
    rates = {row[8]: clean_number(row[9], 0) for row in ws.iter_rows(min_row=8, values_only=True)
             if len(row) > 9 and row[8]}
    products = {row[2]: row[3] for row in ws.iter_rows(min_row=8, values_only=True)
                if len(row) > 3 and row[2] and row[3]}
    rows = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        store, seller, sku = row[0], row[1], row[2]
        if store not in STORES or not seller or not sku:
            continue
        qty = clean_number(row[4], 0)
        accrued = clean_number(row[5], 0)
        rate = rates.get(sku, accrued / qty if qty else 0)
        rows.append({"store": store, "name": seller, "product": row[3] or sku,
                     "sku": sku, "qty": qty, "bonus": accrued,
                     "rate": rate, "expectedBonus": rate * qty})
    rate_details = [{"product": products.get(sku, sku), "sku": sku, "rate": amount}
                    for sku, amount in rates.items()]
    return rows, sorted(rate_details, key=lambda item: item["product"].casefold())


def read_tradein(ws):
    rates = {row[9]: clean_number(row[10], 0) for row in ws.iter_rows(min_row=11, values_only=True)
             if len(row) > 10 and row[9]}
    rows, current_store, current_seller = [], None, None
    for row in ws.iter_rows(min_row=11, values_only=True):
        if row[0]:
            current_store, current_seller = row[0], None
        if row[1]:
            current_seller = row[1]
        if current_store not in STORES or not current_seller or not row[2]:
            continue
        qty = clean_number(row[3], 0)
        accrued = clean_number(row[4], 0)
        rate = rates.get(row[2], 0)
        rows.append({"store": current_store, "name": current_seller, "product": row[2],
                     "qty": qty, "bonus": accrued, "rate": rate,
                     "expectedBonus": rate * qty})
    rate_details = [{"product": product, "rate": amount} for product, amount in rates.items()]
    return rows, sorted(rate_details, key=lambda item: item["product"].casefold())


def read_seller_leaderboard(ws):
    employees = {}
    starts = range(0, min(ws.max_column, 36), 4)
    header = next(ws.iter_rows(min_row=7, max_row=7, values_only=True))
    for row in ws.iter_rows(min_row=9, values_only=True):
        for start in starts:
            name = row[start] if start < len(row) else None
            store = row[start + 1] if start + 1 < len(row) else None
            if not name or store not in STORES:
                continue
            group_label = header[start] if start < len(header) else ""
            employees[name] = {
                "name": name, "store": store,
                "sales": clean_number(row[start + 2], 0),
                "rank": clean_number(row[start + 3], 0),
                "group": group_label,
            }
    return sorted(employees.values(), key=lambda item: (item["rank"] or 999, -item["sales"], item["name"]))


def read_annual(ws):
    wins = {}
    result_columns = [7, 10, 13, 16, 19, 22, 25, 28, 31, 34, 37, 56]
    details = {}
    for row in ws.iter_rows(min_row=6, max_row=22, values_only=True):
        if len(row) > 72 and row[2] in STORES:
            wins[row[2]] = clean_number(row[72], 0)
            store_details = []
            for result_column in result_columns:
                result = row[result_column - 1]
                if result is None or result == "":
                    continue
                start_column = result_column - 2
                title = ws.cell(2, start_column).value
                period = ws.cell(3, start_column).value
                kind = str(ws.cell(4, result_column).value or "").strip().casefold()
                if "мест" in kind:
                    result_label = f"Место {result}"
                elif isinstance(result, (int, float)):
                    result_label = f"Приз {result:,.0f} ₽".replace(",", " ")
                else:
                    result_label = str(result)
                store_details.append({"title": str(title or "Мотивация"), "period": str(period or "Период не указан"), "result": result_label})
            details[row[2]] = store_details
    return wins, details


def main():
    if len(sys.argv) > 1:
        main_book = Path(sys.argv[1]).expanduser()
    else:
        main_book = MAIN_BOOK
    if len(sys.argv) > 2:
        annual_book = Path(sys.argv[2]).expanduser()
    else:
        annual_book = ANNUAL_BOOK
    if not main_book.is_file():
        raise SystemExit(f"Не найден основной файл мотиваций: {main_book}")

    wb = load_workbook(main_book, read_only=True, data_only=True)
    tv_ws = wb["ТВ+саундбары"]
    tab_ws = wb["TabS12+|Ultra"]
    team = {"tv": read_team(tv_ws, True), "tab": read_team(tab_ws, False)}
    ubear_rows, ubear_rates = read_ubear(wb["Доп бонус БЗУ uBear"])
    trade_rows, trade_rates = read_tradein(wb["Trade-in Watch"])
    personal = {"uBear": ubear_rows, "trade": trade_rows}
    employees = read_seller_leaderboard(wb["TabS12+|Ultra продавцы"])
    updated = tv_ws.cell(1, 11).value
    # The campaign period is embedded in cell B5
    # (for example, "Срок проведения: 05.10.2026 - 31.12.2026").
    tv_dates = date_text(tv_ws.cell(5, 2).value)
    tab_dates = date_text(tab_ws.cell(5, 2).value)
    year = updated.year if isinstance(updated, (datetime, date)) else datetime.now().year
    ubear_quarter = wb["Доп бонус БЗУ uBear"].cell(4, 2).value or "Q4"
    trade_quarter = wb["Trade-in Watch"].cell(7, 2).value or "Q4"
    wb.close()

    annual = {name: 0 for name in STORES}
    annual_details = {name: [] for name in STORES}
    if annual_book.is_file():
        annual_wb = load_workbook(annual_book, read_only=True, data_only=True)
        annual, annual_details = read_annual(annual_wb["Рейтинг накопительный"])
        annual_wb.close()

    store_meta = {
        "SMSG_БРН_ГЭЛАКСИ": "БРН", "SMSG_ИРК_МОДНЫЙ_КВАРТАЛ": "ИРК",
        "SMSG_ИРК_СИЛЬВЕР_МОЛЛ": "ИРК", "SMSG_КМР_ЛАПЛАНДИЯ": "КМР",
        "SMSG_НБК_ГАЛЕРЕЯ": "НБК", "SMSG_НБК_КРАСНЫЙ": "НБК", "SMSG_НБК_МЕГА": "НБК",
        "SMSG_НВК_ПЛАНЕТА": "НВК", "SMSG_ОМС_МЕГА": "ОМС", "SMSG_СГТ_АУРА": "СГТ",
        "SMSG_СГТ_СИТИ_МОЛЛ": "СГТ", "SMSG_ТМК_ИЗУМРУДНЫЙ_ГОРОД": "ТМК",
        "SMSG_ТМН_ГУДВИН": "ТМН", "SMSG_ТМН_КРИСТАЛЛ": "ТМН",
        "SMSG_ТМН_СИТИ_МОЛЛ": "ТМН", "SMSG_ХБР_БРОСКО_МОЛЛ": "ХБР",
    }
    stores = []
    for name in STORES:
        tv, tab = team["tv"].get(name), team["tab"].get(name)
        stores.append({
            "name": name, "short": store_meta[name], "closed": name == "SMSG_ИРК_СИЛЬВЕР_МОЛЛ",
            "annualWins": annual.get(name, 0),
            "annualDetails": annual_details.get(name, []),
            "tv": {"sales": tv["fact"], "plan": tv["plan"], "forecast": tv["forecast"]} if tv else {"sales": None, "plan": None, "forecast": None},
            "tab": {"sales": tab["fact"], "plan": tab["plan"], "forecast": tab["forecast"]} if tab else {"sales": None, "plan": None, "forecast": None},
        })

    campaigns = [
        {"id": "tv", "type": "team", "title": "ТВ + саундбары", "sub": "Продажи телевизоров и саундбаров Samsung", "dates": tv_dates or "Период уточняется", "period": short_period(tv_dates, "ТВ"), "conditions": "Период: 05.10–31.12.2026. Участвуют магазины и региональные директора; учитываются все каналы продаж. Каждый магазин относится к одной из 12 групп и получает призовое место только при выполнении плана выше 100%. Призовой фонд распределяет внутри магазина управляющий. План региона равен сумме планов магазинов. Региональные директора соревнуются по перевыполнению плана; призы начисляются в премию за декабрь.", "style": ""},
        {"id": "tab", "type": "team", "title": "Планшеты Tab S12", "sub": "Продажи планшетов Samsung Tab S12", "dates": tab_dates or "Период уточняется", "period": short_period(tab_dates, "Tab S12"), "conditions": "Период: 19.10–30.11.2026. В зачёт идут Galaxy Tab S12+ и Tab S12 Ultra всех конфигураций, проданные офлайн или с самовывозом. Это три независимых зачёта: магазин получает Galaxy Buds4 при выполнении собственного плана на 100%; продавцы соревнуются по личным продажам внутри группы магазинов, а Galaxy Buds4 Pro получают первые места при достижении порога группы; региональный директор получает Galaxy Buds4 Pro, если продажи магазинов региона не ниже суммы их планов. Выполнение цели сети на выдачу призов не влияет. Призы магазина и региона распределяются по правилам мотивации.", "style": "tab"},
        {"id": "trade", "type": "personal", "title": "Trade-in Watch", "sub": "Личное вознаграждение за продажу часов", "dates": f"02.10–31.10.{year}", "period": f"02.10–31.10", "conditions": "Период: 02.10–31.10.2026. За каждую Galaxy Watch9 или Galaxy Watch Ultra 2, проданную в trade-in, продавцу начисляется бонус по модели: 800 ₽ за Watch Ultra 2 и 500 ₽ за Watch9. Учитываются продажи офлайн и заказы с самовывозом. Обязательное условие выплат — выполнение сетью плана 230 часов за период; при невыполнении плана бонус не выплачивается. Выплата начисляется в премию.", "rates": trade_rates, "style": "trade-card"},
        {"id": "uBear", "type": "personal", "title": "БЗУ uBear", "sub": "Личный бонус за продажу зарядных устройств", "dates": f"01.10–30.11.{year}", "period": "01.10–30.11", "conditions": "Период: 01.10–30.11.2026. Продавец получает бонус за каждое беспроводное зарядное устройство uBear из списка — без плана и дополнительных условий. Учитываются продажи офлайн и заказы с самовывозом. Размер бонуса зависит от модели; ставки указаны ниже. Выплата начисляется в премию.", "rates": ubear_rates, "style": "ubear-card"},
    ]
    payload = {
        "updatedAt": updated.isoformat() if isinstance(updated, (datetime, date)) else str(updated or ""),
        "stores": stores, "teamStats": team, "personalResults": personal,
        "employees": employees, "campaigns": campaigns,
    }
    OUT.write_text("window.SITE_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(f"Сформирован {OUT.name}: {len(stores)} магазинов, {len(campaigns)} мотивации, дата среза {payload['updatedAt']}")


if __name__ == "__main__":
    main()
